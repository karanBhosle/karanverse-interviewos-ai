from io import BytesIO

import pytest
from docx import Document


def _resume_file() -> tuple[str, bytes, str]:
    document = Document()
    document.add_paragraph("Karan Bhosle")
    document.add_paragraph("Skills")
    document.add_paragraph("Python, SQL, retrieval augmented generation, LangGraph")
    document.add_paragraph("Experience")
    document.add_paragraph("Data Scientist – ExampleCo – 2024")
    document.add_paragraph("- Built a FastAPI RAG service and improved recall 12%")
    buffer = BytesIO()
    document.save(buffer)
    return "resume.docx", buffer.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


@pytest.mark.asyncio
async def test_local_workspace_without_token(client):
    from app.core.config import get_settings

    get_settings().auth_required = False
    try:
        response = await client.get("/api/v1/auth/me")
        assert response.status_code == 200
        assert response.json()["email"] == "karan@local"
    finally:
        get_settings().auth_required = True


@pytest.mark.asyncio
async def test_health(client):
    response = await client.get("/api/v1/health")
    assert response.status_code == 200
    assert response.json()["status"] == "ok"


@pytest.mark.asyncio
async def test_auth_and_ownership(client, auth_client):
    other = await client.post(
        "/api/v1/auth/register",
        json={"email": "other@example.com", "password": "password12", "display_name": "Other"},
    )
    other_token = other.json()["access_token"]
    filename, data, content_type = _resume_file()
    uploaded = await auth_client.post(
        "/api/v1/resumes",
        files={"file": (filename, data, content_type)},
    )
    assert uploaded.status_code == 200
    resume_id = uploaded.json()["id"]
    forbidden = await client.get(
        f"/api/v1/resumes/{resume_id}",
        headers={"Authorization": f"Bearer {other_token}"},
    )
    assert forbidden.status_code == 403


@pytest.mark.asyncio
async def test_text_resume_and_job_create(auth_client):
    uploaded = await auth_client.post(
        "/api/v1/resumes",
        files={"file": ("resume.txt", b"Karan Bhosle\nSkills\nPython, SQL\nExperience\n- Built a RAG service", "text/plain")},
    )
    assert uploaded.status_code == 200
    job = await auth_client.post(
        "/api/v1/jobs",
        json={"title": "Data Scientist", "raw_text": "Need Python, SQL, and RAG experience."},
    )
    assert job.status_code == 200
    assert job.json()["requirements"]
    too_short = await auth_client.post("/api/v1/jobs", json={"title": "DS", "raw_text": "Hi"})
    assert too_short.status_code == 422


@pytest.mark.asyncio
async def test_upload_validation(auth_client):
    response = await auth_client.post(
        "/api/v1/resumes",
        files={"file": ("notes.exe", b"MZ", "application/octet-stream")},
    )
    assert response.status_code == 422


@pytest.mark.asyncio
async def test_session_persistence_and_failed_llm(auth_client):
    filename, data, content_type = _resume_file()
    resume = await auth_client.post("/api/v1/resumes", files={"file": (filename, data, content_type)})
    job = await auth_client.post(
        "/api/v1/jobs",
        json={
            "title": "ML Engineer",
            "company": "Example",
            "raw_text": "Requirements:\n- Strong Python experience\n- retrieval augmented generation\n- LangGraph",
        },
    )
    match = await auth_client.post(
        "/api/v1/job-matches",
        json={"resume_id": resume.json()["id"], "job_id": job.json()["id"]},
    )
    assert match.status_code == 200
    assert "not a probability" in match.json()["score_explanation"].lower() or "not a probability of being hired" in match.json()["score_explanation"]
    plan = await auth_client.post(
        "/api/v1/interview-plans",
        json={
            "resume_id": resume.json()["id"],
            "job_id": job.json()["id"],
            "mode": "mixed",
            "difficulty": "medium",
            "duration_minutes": 20,
        },
    )
    session = await auth_client.post(
        "/api/v1/interview-sessions",
        json={"plan_id": plan.json()["id"], "visual_consent": False},
    )
    assert session.status_code == 200
    assert session.json()["speech_assisted"] is True
    question = session.json()["current_question"]
    assert question is not None
    answered = await auth_client.post(
        f"/api/v1/interview-sessions/{session.json()['id']}/answers",
        json={
            "question_id": question["id"],
            "raw_text": (
                "I would treat this as a typed state machine with deterministic scoring, "
                "iteration limits, and fail-closed evaluation when the model is unavailable."
            ),
            "transcript_provenance": "speech_to_text",
        },
    )
    assert answered.status_code == 200
    assert answered.json()["evaluation"]["provenance"] == "deterministic_concept_coverage"
    completed = await auth_client.post(f"/api/v1/interview-sessions/{session.json()['id']}/complete")
    assert completed.status_code == 200
    report = await auth_client.get(f"/api/v1/interview-sessions/{session.json()['id']}/report")
    assert report.status_code == 200
    assert report.json()["answers"]
    assert report.json()["communication_feedback"][0]["transcript_provenance"] == "speech_to_text"


@pytest.mark.asyncio
async def test_data_deletion(auth_client):
    filename, data, content_type = _resume_file()
    await auth_client.post("/api/v1/resumes", files={"file": (filename, data, content_type)})
    deleted = await auth_client.delete("/api/v1/users/me/data")
    assert deleted.status_code == 200
    remaining = await auth_client.get("/api/v1/resumes")
    assert remaining.json() == []


@pytest.mark.asyncio
async def test_visual_feedback_requires_consent(auth_client):
    plan = await auth_client.post(
        "/api/v1/interview-plans",
        json={"mode": "behavioral", "difficulty": "medium", "duration_minutes": 15},
    )
    session = await auth_client.post(
        "/api/v1/interview-sessions",
        json={"plan_id": plan.json()["id"], "visual_consent": False},
    )
    ingest = await auth_client.post(
        f"/api/v1/interview-sessions/{session.json()['id']}/visual-feedback",
        json={"consent": True, "frames": []},
    )
    assert ingest.json()["accepted"] is False


@pytest.mark.asyncio
async def test_tts_returns_neural_audio(auth_client, monkeypatch):
    async def fake_indian(text, voice, rate, pitch):
        assert "Python" in text
        assert voice == "en-IN-NeerjaNeural"
        return b"ID3FAKEAUDIO", "audio/mpeg"

    monkeypatch.setattr("app.services.tts._neural_indian", fake_indian)
    empty = await auth_client.post("/api/v1/tts", json={"text": ""})
    assert empty.status_code == 422
    spoken = await auth_client.post("/api/v1/tts", json={"text": "Explain Python lists and tuples."})
    assert spoken.status_code == 200
    assert spoken.headers["content-type"].startswith("audio/")
    assert spoken.headers["x-interviewos-voice"] == "en-IN-NeerjaNeural"
    assert spoken.content.startswith(b"ID3")
