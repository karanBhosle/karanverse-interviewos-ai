from __future__ import annotations

import io
import logging
from dataclasses import dataclass

from pypdf import PdfReader
from docx import Document

from app.core.errors import ValidationAppError

logger = logging.getLogger(__name__)

ALLOWED_TYPES = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "text/plain": "text",
}


@dataclass(frozen=True)
class ExtractedPage:
    page: int
    text: str


@dataclass(frozen=True)
class ExtractedDocument:
    text: str
    pages: list[ExtractedPage]
    ocr_used: bool
    method: str


def sniff_kind(filename: str, content_type: str | None) -> str:
    lowered = filename.lower()
    if lowered.endswith(".pdf"):
        return "pdf"
    if lowered.endswith(".docx"):
        return "docx"
    if lowered.endswith(".txt"):
        return "text"
    if lowered.endswith(".doc"):
        raise ValidationAppError("Old .doc files are not supported. Save the resume as PDF or DOCX, or paste the text.")
    if content_type in ALLOWED_TYPES:
        return ALLOWED_TYPES[content_type]
    raise ValidationAppError("Only PDF, DOCX, or plain-text resumes are accepted", {"filename": filename})


def extract_text(data: bytes, kind: str, *, enable_ocr: bool = False) -> ExtractedDocument:
    if kind == "pdf":
        return _extract_pdf(data, enable_ocr=enable_ocr)
    if kind == "docx":
        return _extract_docx(data)
    if kind == "text":
        return _extract_plain_text(data)
    raise ValidationAppError("Unsupported file type")


def _extract_plain_text(data: bytes) -> ExtractedDocument:
    text = data.decode("utf-8", errors="replace").strip()
    return ExtractedDocument(
        text=text,
        pages=[ExtractedPage(page=1, text=text)] if text else [],
        ocr_used=False,
        method="plain_text",
    )


def _extract_pdf(data: bytes, *, enable_ocr: bool) -> ExtractedDocument:
    try:
        reader = PdfReader(io.BytesIO(data), strict=False)
    except Exception as exc:
        raise ValidationAppError("The PDF could not be read. The file may be damaged.") from exc

    pages: list[ExtractedPage] = []
    for index, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        pages.append(ExtractedPage(page=index, text=text.strip()))

    joined = "\n\n".join(page.text for page in pages if page.text)
    ocr_used = False
    if not joined.strip() and enable_ocr:
        joined, ocr_used = _ocr_pdf_fallback(data)
        pages = [ExtractedPage(page=1, text=joined)] if joined else pages
    return ExtractedDocument(text=joined, pages=pages, ocr_used=ocr_used, method="pypdf")


def _extract_docx(data: bytes) -> ExtractedDocument:
    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:
        raise ValidationAppError("The DOCX file could not be read.") from exc
    paragraphs = [p.text.strip() for p in document.paragraphs if p.text and p.text.strip()]
    text = "\n".join(paragraphs)
    return ExtractedDocument(
        text=text,
        pages=[ExtractedPage(page=1, text=text)],
        ocr_used=False,
        method="python-docx",
    )


def _ocr_pdf_fallback(data: bytes) -> tuple[str, bool]:
    try:
        import pytesseract  # type: ignore
        from pdf2image import convert_from_bytes  # type: ignore
    except Exception:
        logger.info("OCR fallback requested but optional OCR dependencies are not installed")
        return "", False
    try:
        images = convert_from_bytes(data, fmt="png")
        texts = [pytesseract.image_to_string(image) for image in images]
        return "\n\n".join(t.strip() for t in texts if t.strip()), True
    except Exception:
        logger.exception("OCR fallback failed")
        return "", False
